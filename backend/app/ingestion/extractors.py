"""Payload-to-text extractors, shared by every loader.

A URL fetch and a file upload differ only in how the bytes arrive, so the
extraction logic lives here once. `extract_payload` dispatches on content
type and file extension and returns the text plus the source kind it
detected. Each extractor is synchronous and CPU-bound; callers run them in
a worker thread.
"""

from __future__ import annotations

import io
import logging
import re
from typing import Any

from app.core.exceptions import IngestionError
from app.domain.models import SourceKind

log = logging.getLogger(__name__)

_MULTI_BLANK = re.compile(r"\n{3,}")
_SPACES = re.compile(r"[ \t ]+")
_NOISE_TAGS = (
    "script", "style", "noscript", "svg", "iframe", "header", "footer", "nav", "form", "aside",
)

AUDIO_EXTENSIONS = (".mp3", ".m4a", ".wav", ".ogg", ".oga", ".opus", ".flac", ".aac", ".webm", ".mp4")

_EXTENSION_KINDS: dict[str, SourceKind] = {
    ".pdf": SourceKind.PDF,
    ".docx": SourceKind.DOCX,
    ".pptx": SourceKind.PPTX,
    ".xlsx": SourceKind.XLSX,
    ".xlsm": SourceKind.XLSX,
    ".epub": SourceKind.EPUB,
    ".txt": SourceKind.TEXT,
    ".md": SourceKind.TEXT,
    ".markdown": SourceKind.TEXT,
    ".csv": SourceKind.TEXT,
    ".json": SourceKind.TEXT,
    ".rst": SourceKind.TEXT,
    ".htm": SourceKind.WEB,
    ".html": SourceKind.WEB,
}

_MIME_KINDS: dict[str, SourceKind] = {
    "application/pdf": SourceKind.PDF,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": SourceKind.DOCX,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": SourceKind.PPTX,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": SourceKind.XLSX,
    "application/epub+zip": SourceKind.EPUB,
    "text/html": SourceKind.WEB,
    "application/xhtml+xml": SourceKind.WEB,
    "text/plain": SourceKind.TEXT,
    "text/markdown": SourceKind.TEXT,
    "text/csv": SourceKind.TEXT,
    "application/json": SourceKind.TEXT,
}


def tidy(text: str) -> str:
    text = _SPACES.sub(" ", text)
    lines = [line.strip() for line in text.splitlines()]
    return _MULTI_BLANK.sub("\n\n", "\n".join(lines)).strip()


def _filename_of(reference: str) -> str:
    return reference.split("?")[0].split("#")[0].rstrip("/").rsplit("/", 1)[-1]


def _extension_of(reference: str) -> str:
    name = _filename_of(reference)
    return name[name.rfind(".") :].lower() if "." in name else ""


def detect_kind(content_type: str, reference: str) -> SourceKind:
    """Best guess at the kind, preferring an unambiguous file extension."""
    extension = _extension_of(reference)
    if extension in AUDIO_EXTENSIONS and extension != ".mp4":
        return SourceKind.AUDIO
    if extension in _EXTENSION_KINDS:
        return _EXTENSION_KINDS[extension]
    mime = (content_type or "").split(";")[0].strip().lower()
    if mime.startswith("audio/"):
        return SourceKind.AUDIO
    if mime in _MIME_KINDS:
        return _MIME_KINDS[mime]
    return SourceKind.WEB


# --------------------------------------------------------------------------
# format-specific extractors
# --------------------------------------------------------------------------


def extract_html(html: str, url: str) -> tuple[str, str, dict[str, Any]]:
    import trafilatura
    from bs4 import BeautifulSoup

    meta: dict[str, Any] = {}
    title = ""
    try:
        extracted = trafilatura.extract_metadata(html, default_url=url)
        if extracted is not None:
            title = extracted.title or ""
            meta = {
                "description": extracted.description,
                "site_name": extracted.sitename,
                "author": extracted.author,
                "published": extracted.date,
            }
    except Exception as exc:  # noqa: BLE001 - metadata is optional
        log.debug("metadata extraction failed for %s: %s", url, exc)

    text = (
        trafilatura.extract(
            html,
            url=url,
            include_tables=True,
            include_comments=False,
            include_links=False,
            favor_recall=True,
            output_format="txt",
        )
        or ""
    )

    if len(text) < 400:  # readability gave up: fall back to a plain DOM strip
        soup = BeautifulSoup(html, "lxml")
        for tag in soup.find_all(_NOISE_TAGS):
            tag.decompose()
        if not title and soup.title and soup.title.string:
            title = soup.title.string
        body = soup.body or soup
        fallback = tidy(body.get_text("\n"))
        if len(fallback) > len(text):
            text = fallback
            meta["extractor"] = "dom"
    else:
        meta["extractor"] = "trafilatura"

    if not title:
        soup_title = BeautifulSoup(html, "lxml").title
        title = (soup_title.string or "").strip() if soup_title else ""
    title = " ".join((title or url).split())[:300]
    return title, tidy(text), {k: v for k, v in meta.items() if v}


def extract_pdf(payload: bytes, reference: str) -> tuple[str, str, dict[str, Any]]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(payload))
    except Exception as exc:  # noqa: BLE001
        raise IngestionError("The PDF could not be parsed.") from exc

    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - skip an unreadable page
            pages.append("")
    text = tidy("\n\n".join(pages))
    info = reader.metadata or {}
    title = (info.get("/Title") or "").strip() or _filename_of(reference) or "PDF document"
    author = (info.get("/Author") or "").strip() or None
    return title, text, {"pages": len(reader.pages), "author": author}


def extract_docx(payload: bytes, reference: str) -> tuple[str, str, dict[str, Any]]:
    import docx

    try:
        document = docx.Document(io.BytesIO(payload))
    except Exception as exc:  # noqa: BLE001
        raise IngestionError("The Word document could not be parsed.") from exc

    blocks: list[str] = []
    for paragraph in document.paragraphs:
        content = paragraph.text.strip()
        if not content:
            continue
        style = (paragraph.style.name if paragraph.style else "") or ""
        blocks.append(f"\n## {content}\n" if style.startswith("Heading") else content)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))

    props = document.core_properties
    title = (props.title or "").strip() or _filename_of(reference) or "Word document"
    meta = {
        "author": (props.author or "").strip() or None,
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
    }
    return title, tidy("\n".join(blocks)), meta


def extract_pptx(payload: bytes, reference: str) -> tuple[str, str, dict[str, Any]]:
    from pptx import Presentation

    try:
        deck = Presentation(io.BytesIO(payload))
    except Exception as exc:  # noqa: BLE001
        raise IngestionError("The PowerPoint file could not be parsed.") from exc

    blocks: list[str] = []
    for number, slide in enumerate(deck.slides, start=1):
        lines = [f"\n## Slide {number}"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    content = "".join(run.text for run in paragraph.runs).strip()
                    if content:
                        lines.append(content)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        if notes:
            lines.append(f"Speaker notes: {notes}")
        if len(lines) > 1:
            blocks.append("\n".join(lines))

    title = (deck.core_properties.title or "").strip() or _filename_of(reference) or "Presentation"
    return title, tidy("\n".join(blocks)), {"slides": len(deck.slides)}


def extract_xlsx(payload: bytes, reference: str) -> tuple[str, str, dict[str, Any]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise IngestionError("The spreadsheet could not be parsed.") from exc

    blocks: list[str] = []
    rows_total = 0
    try:
        for sheet in workbook.worksheets:
            lines = [f"\n## Sheet: {sheet.title}"]
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    lines.append(" | ".join(cells))
                    rows_total += 1
            if len(lines) > 1:
                blocks.append("\n".join(lines))
    finally:
        workbook.close()

    title = _filename_of(reference) or "Spreadsheet"
    return title, tidy("\n".join(blocks)), {"sheets": len(blocks), "rows": rows_total}


def extract_epub(payload: bytes, reference: str) -> tuple[str, str, dict[str, Any]]:
    import tempfile
    from pathlib import Path

    import ebooklib
    from bs4 import BeautifulSoup
    from ebooklib import epub

    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "book.epub"
        path.write_bytes(payload)
        try:
            book = epub.read_epub(str(path), options={"ignore_ncx": True})
        except Exception as exc:  # noqa: BLE001
            raise IngestionError("The EPUB file could not be parsed.") from exc

        chapters: list[str] = []
        for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
            soup = BeautifulSoup(item.get_content(), "lxml")
            for tag in soup.find_all(_NOISE_TAGS):
                tag.decompose()
            content = tidy(soup.get_text("\n"))
            if len(content) > 40:
                chapters.append(content)

        def first(name: str) -> str | None:
            values = book.get_metadata("DC", name)
            return str(values[0][0]) if values else None

        title = (first("title") or "").strip() or _filename_of(reference) or "EPUB book"
        meta = {"author": first("creator"), "language": first("language"), "chapters": len(chapters)}
    return title, tidy("\n\n".join(chapters)), meta


def extract_text(payload: bytes | str, reference: str) -> tuple[str, str, dict[str, Any]]:
    raw = payload.decode("utf-8", "replace") if isinstance(payload, bytes) else payload
    return _filename_of(reference) or reference, tidy(raw), {}


# --------------------------------------------------------------------------
# dispatcher
# --------------------------------------------------------------------------

_BINARY_EXTRACTORS = {
    SourceKind.PDF: extract_pdf,
    SourceKind.DOCX: extract_docx,
    SourceKind.PPTX: extract_pptx,
    SourceKind.XLSX: extract_xlsx,
    SourceKind.EPUB: extract_epub,
}

SUPPORTED_UPLOAD_EXTENSIONS = tuple(
    sorted({*_EXTENSION_KINDS, *AUDIO_EXTENSIONS})
)


def extract_payload(
    payload: bytes, *, content_type: str, reference: str
) -> tuple[str, str, dict[str, Any], SourceKind]:
    """Return (title, text, metadata, kind) for a fetched or uploaded payload.

    Audio is not handled here: it needs a transcription provider and produces
    timed segments, so `AudioLoader` owns it.
    """
    kind = detect_kind(content_type, reference)
    if kind is SourceKind.AUDIO:
        raise IngestionError("Audio payloads are transcribed by the audio loader, not extracted.")

    extractor = _BINARY_EXTRACTORS.get(kind)
    if extractor is not None:
        title, text, meta = extractor(payload, reference)
    elif kind is SourceKind.TEXT:
        title, text, meta = extract_text(payload, reference)
    else:
        html = payload.decode("utf-8", "replace") if isinstance(payload, bytes) else payload
        title, text, meta = extract_html(html, reference)
        kind = SourceKind.WEB
    return title, text, meta, kind
