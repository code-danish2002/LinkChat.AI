"""Extractor round-trips against real files built by the same libraries."""

from __future__ import annotations

import io

import pytest

from app.core.exceptions import IngestionError
from app.domain.models import SourceKind
from app.ingestion.extractors import detect_kind, extract_payload, tidy


def build_docx() -> bytes:
    import docx

    document = docx.Document()
    document.core_properties.title = "Quarterly Report"
    document.core_properties.author = "Platform Team"
    document.add_heading("Revenue", level=1)
    document.add_paragraph("Revenue reached 4.2 million euros in the quarter.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Share"
    table.cell(1, 0).text = "EMEA"
    table.cell(1, 1).text = "38%"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pptx() -> bytes:
    from pptx import Presentation

    deck = Presentation()
    deck.core_properties.title = "Launch Plan"
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Timeline"
    slide.placeholders[1].text = "Beta opens on 3 March"
    slide.notes_slide.notes_text_frame.text = "Mention the pricing change"
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def build_xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    sheet.append(["SKU", "Price"])
    sheet.append(["X-9000", 49])
    sheet.append(["X-9100", 79])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_epub() -> bytes:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("id1")
    book.set_title("A Short Book")
    book.set_language("en")
    book.add_author("Some Author")
    chapter = epub.EpubHtml(title="One", file_name="one.xhtml", lang="en")
    chapter.content = "<h1>Chapter One</h1><p>" + "The sky was a bright shade of orange. " * 5 + "</p>"
    book.add_item(chapter)
    book.spine = ["nav", chapter]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    buffer = io.BytesIO()
    epub.write_epub(buffer, book)
    return buffer.getvalue()


@pytest.mark.parametrize(
    "content_type,reference,expected",
    [
        ("application/pdf", "https://x.test/paper", SourceKind.PDF),
        ("", "https://x.test/paper.pdf", SourceKind.PDF),
        ("", "report.docx", SourceKind.DOCX),
        ("", "deck.pptx", SourceKind.PPTX),
        ("", "sheet.xlsx", SourceKind.XLSX),
        ("", "book.epub", SourceKind.EPUB),
        ("", "talk.mp3", SourceKind.AUDIO),
        ("audio/wav", "stream", SourceKind.AUDIO),
        ("text/plain", "notes", SourceKind.TEXT),
        ("", "notes.md", SourceKind.TEXT),
        ("text/html", "https://x.test/page", SourceKind.WEB),
        ("", "https://x.test/page", SourceKind.WEB),
    ],
)
def test_detect_kind(content_type, reference, expected):
    assert detect_kind(content_type, reference) is expected


def test_docx_round_trip():
    title, text, meta, kind = extract_payload(build_docx(), content_type="", reference="report.docx")
    assert kind is SourceKind.DOCX
    assert title == "Quarterly Report"
    assert meta["author"] == "Platform Team"
    assert "## Revenue" in text
    assert "4.2 million euros" in text
    assert "EMEA | 38%" in text


def test_pptx_round_trip():
    title, text, meta, kind = extract_payload(build_pptx(), content_type="", reference="deck.pptx")
    assert kind is SourceKind.PPTX
    assert title == "Launch Plan"
    assert meta["slides"] == 1
    assert "## Slide 1" in text
    assert "Beta opens on 3 March" in text
    assert "Mention the pricing change" in text


def test_xlsx_round_trip():
    title, text, meta, kind = extract_payload(build_xlsx(), content_type="", reference="sheet.xlsx")
    assert kind is SourceKind.XLSX
    assert title == "sheet.xlsx"
    assert meta["rows"] == 3
    assert "## Sheet: Prices" in text
    assert "X-9000 | 49" in text


def test_epub_round_trip():
    title, text, meta, kind = extract_payload(build_epub(), content_type="", reference="book.epub")
    assert kind is SourceKind.EPUB
    assert title == "A Short Book"
    assert meta["author"] == "Some Author"
    assert "bright shade of orange" in text


def test_html_extraction_prefers_article_text():
    html = """
    <html><head><title>Ignored chrome</title></head><body>
      <nav>home about contact</nav>
      <article><h1>Hybrid search</h1>
        <p>%s</p>
        <p>Keyword scoring finds exact product codes such as X-9000 that vectors miss.</p>
      </article>
      <footer>copyright</footer>
    </body></html>
    """ % ("Combining dense vectors with keyword scoring improves recall on short queries. " * 8)
    _, text, _, kind = extract_payload(html.encode(), content_type="text/html", reference="https://x.test/a")
    assert kind is SourceKind.WEB
    assert "X-9000" in text
    assert "copyright" not in text.lower()


def test_corrupt_office_file_raises_a_clear_error():
    with pytest.raises(IngestionError, match="Word document"):
        extract_payload(b"not a docx at all", content_type="", reference="broken.docx")


def test_audio_is_not_handled_by_the_text_extractors():
    with pytest.raises(IngestionError, match="audio loader"):
        extract_payload(b"\x00\x01", content_type="audio/mpeg", reference="a.mp3")


def test_tidy_collapses_whitespace():
    assert tidy("a  \t b\n\n\n\nc  ") == "a b\n\nc"
